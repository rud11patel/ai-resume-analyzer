import io
from typing import Dict, Any, List
import pymupdf as fitz
from docx import Document


class DocumentExtractionError(Exception):
    """Custom exception raised when resume text extraction encounters an error."""
    def __init__(self, message: str, status_code: int = 400):
        super().__init__(message)
        self.message = message
        self.status_code = status_code


def extract_text_from_pdf(content: bytes) -> Dict[str, Any]:
    """
    Extracts plain text and page metadata from a PDF byte buffer using PyMuPDF (fitz).
    Handles empty documents, encrypted PDFs, and scanned PDFs with no text layer.
    """
    if not content or len(content) == 0:
        raise DocumentExtractionError("The uploaded PDF file is empty (0 bytes).")

    try:
        doc = fitz.open(stream=content, filetype="pdf")
    except Exception as exc:
        raise DocumentExtractionError(f"Failed to parse PDF document. File may be corrupt or invalid: {str(exc)}")

    try:
        if doc.is_encrypted:
            raise DocumentExtractionError("Encrypted or password-protected PDFs are not supported.", status_code=422)

        page_count = len(doc)
        if page_count == 0:
            raise DocumentExtractionError("PDF contains 0 pages.")

        extracted_pages: List[str] = []
        for page_idx in range(page_count):
            page = doc.load_page(page_idx)
            text = page.get_text("text") or ""
            extracted_pages.append(text.strip())

        full_text = "\n\n".join(filter(None, extracted_pages)).strip()
        warnings: List[str] = []

        # Scanned document detection heuristic:
        # If there are pages but virtually no text (< 30 characters total across pages),
        # the document is likely an image-only scanned PDF without an OCR text layer.
        if len(full_text) < 30:
            warnings.append(
                "Document appears to be a scanned image with minimal or no extractable text layer. "
                "For optimal ATS parsing, please upload a digital PDF with selectable text."
            )

        char_count = len(full_text)
        word_count = len(full_text.split()) if char_count > 0 else 0

        return {
            "success": True,
            "text": full_text,
            "metadata": {
                "characterCount": char_count,
                "wordCount": word_count,
                "pageCount": page_count,
                "fileType": "pdf",
            },
            "warnings": warnings,
        }
    finally:
        doc.close()


def extract_text_from_docx(content: bytes) -> Dict[str, Any]:
    """
    Extracts text from paragraphs and tables in a Microsoft Word DOCX byte buffer using python-docx.
    """
    if not content or len(content) == 0:
        raise DocumentExtractionError("The uploaded DOCX file is empty (0 bytes).")

    try:
        buffer = io.BytesIO(content)
        doc = Document(buffer)
    except Exception as exc:
        raise DocumentExtractionError(f"Failed to parse DOCX document. File may be corrupt or malformed: {str(exc)}")

    paragraphs_text = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            paragraphs_text.append(t)

    # Also extract text inside tables (resumes often format skills/experience in tables)
    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                paragraphs_text.append(" | ".join(row_cells))

    full_text = "\n\n".join(paragraphs_text).strip()
    warnings: List[str] = []

    if len(full_text) < 20:
        warnings.append(
            "DOCX file contains very little text. Please ensure the document is not empty."
        )

    char_count = len(full_text)
    word_count = len(full_text.split()) if char_count > 0 else 0

    return {
        "success": True,
        "text": full_text,
        "metadata": {
            "characterCount": char_count,
            "wordCount": word_count,
            "pageCount": 1,  # DOCX format does not have fixed physical pages until rendered
            "fileType": "docx",
        },
        "warnings": warnings,
    }


def extract_document_text(filename: str, content: bytes) -> Dict[str, Any]:
    """
    Dispatches document extraction based on filename extension.
    """
    lower_name = (filename or "").lower().strip()
    if lower_name.endswith(".pdf"):
        return extract_text_from_pdf(content)
    elif lower_name.endswith(".docx"):
        return extract_text_from_docx(content)
    else:
        raise DocumentExtractionError(
            f"Unsupported document format for '{filename}'. Only PDF (.pdf) and DOCX (.docx) files are supported."
        )
