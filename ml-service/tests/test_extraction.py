import io
import pytest
from fastapi.testclient import TestClient
import pymupdf as fitz
from docx import Document

from app.main import app

client = TestClient(app)


def test_health_endpoint():
    response = client.get("/health")
    assert response.status_code == 200
    data = response.json()
    assert data["status"] == "ok"
    assert "pdf-extraction" in data["capabilities"]


def test_reject_unsupported_extension():
    response = client.post(
        "/extract-text",
        files={"file": ("malicious.exe", b"binarycontent", "application/octet-stream")},
    )
    assert response.status_code == 400
    data = response.json()
    assert data["success"] is False
    assert "Unsupported file format" in data["error"]


def test_reject_empty_file():
    response = client.post(
        "/extract-text",
        files={"file": ("empty_resume.pdf", b"", "application/pdf")},
    )
    assert response.status_code == 400
    data = response.json()
    assert data["success"] is False
    assert "empty" in data["error"].lower()


def test_extract_pdf_success():
    # Generate a clean in-memory PDF using PyMuPDF
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((50, 72), "Alex Morgan\nSenior Full Stack Engineer\nSkills: React, Node.js, Python, MongoDB")
    pdf_bytes = doc.tobytes()
    doc.close()

    response = client.post(
        "/extract-text",
        files={"file": ("test_resume.pdf", pdf_bytes, "application/pdf")},
    )
    assert response.status_code == 200
    data = response.json()
    assert data["success"] is True
    assert "Alex Morgan" in data["text"]
    assert "Skills: React" in data["text"]
    assert data["metadata"]["fileType"] == "pdf"
    assert data["metadata"]["pageCount"] == 1
    assert data["metadata"]["characterCount"] > 20


def test_extract_docx_success():
    # Generate a clean in-memory DOCX using python-docx
    doc = Document()
    doc.add_heading("Candidate Resume", 0)
    doc.add_paragraph("Experienced Cloud Architect with 8+ years designing scalable systems.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Core Competencies"
    table.rows[0].cells[1].text = "AWS, Docker, Kubernetes"

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    response = client.post(
        "/extract-text",
        files={"file": ("test_resume.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 200
    data = response.json()
    assert data["success"] is True
    assert "Experienced Cloud Architect" in data["text"]
    assert "AWS, Docker, Kubernetes" in data["text"]
    assert data["metadata"]["fileType"] == "docx"
    assert data["metadata"]["characterCount"] > 30


def test_corrupt_pdf_handling():
    corrupt_bytes = b"%PDF-1.4\ncorrupt header and unreadable garbage payload without endobj"
    response = client.post(
        "/extract-text",
        files={"file": ("broken.pdf", corrupt_bytes, "application/pdf")},
    )
    assert response.status_code in (400, 422, 500)
    data = response.json()
    assert data["success"] is False
